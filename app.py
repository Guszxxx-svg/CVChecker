import streamlit as st
import openpyxl
import time
from google import genai

# Konfigurasi Halaman Web
st.set_page_config(page_title="AI CV Checker & Jouken Matcher", page_icon="📄", layout="wide")

st.title("📄 AI CV Checker & Jouken Matcher")
st.write("Unggah file Excel daftar kandidat multi-sheet Anda dan tentukan syarat (*Jouken*) untuk disaring secara otomatis oleh AI.")

# Input API Key (bisa diisi otomatis atau dimasukkan pengguna)
api_key_default = "AQ.Ab8RN6Lvb9zmxhY2tevySxCNtHls66YB_1y5qscEhhA7AdNy1g"
api_key = st.text_input("Gemini API Key:", value=api_key_default, type="password")

# Input Syarat / Jouken
jouken_input = st.text_area(
    "Masukkan Syarat / Jouken Klien (Contoh: Tokyo, gaji 13万円, tidak boleh makan babi, dll):",
    height=100
)

# Upload File Excel
uploaded_file = st.file_uploader("Pilih file Excel Kandidat (.xlsx)", type=["xlsx"])

if st.button("Mulai Proses Pengecekan CV", type="primary"):
    if not api_key:
        st.error("Mohon masukkan Gemini API Key terlebih dahulu!")
    elif not jouken_input:
        st.error("Mohon masukkan Syarat / Jouken terlebih dahulu!")
    elif uploaded_file is None:
        st.error("Mohon unggah file Excel kandidat terlebih dahulu!")
    else:
        st.info("Memproses file Excel dan menganalisis kandidat...")
        
        try:
            # Load workbook langsung dari file yang di-upload
            wb = openpyxl.load_workbook(uploaded_file, data_only=True)
            sheet_names = wb.sheetnames
            
            client = genai.Client(api_key=api_key)
            
            total_candidates = len(sheet_names) - 1
            st.write(f"Total kandidat (sheet) ditemukan: {total_candidates}")
            
            progress_bar = st.progress(0)
            
            # Loop setiap sheet kandidat
            for i in range(1, len(sheet_names)):
                current_sheet_name = sheet_names[i]
                
                # Lewati sheet template/sistem
                if "sheet" in current_sheet_name.lower() and len(current_sheet_name) < 10:
                    continue
                
                ws = wb[current_sheet_name]
                cv_text_lines = []
                for row in ws.iter_rows(values_only=True):
                    row_str = " ".join([str(cell) for cell in row if cell is not None])
                    if row_str.strip():
                        cv_text_lines.append(row_str)
                
                full_cv_text = "\n".join(cv_text_lines)
                
                if not full_cv_text.strip():
                    continue
                
                # Buat Prompt untuk Gemini
                prompt = f"""
                Bertindaklah sebagai Professional HR Recruiter. Analisis CV kandidat ini berdasarkan SYARAT KETAT (Jouken) yang diberikan.
                
                SYARAT KETAT (Jouken) DARI KLIEN:
                {jouken_input}
                
                TEKS CV KANDIDAT ({current_sheet_name}):
                {full_cv_text}
                
                TUGAS:
                Periksa apakah kandidat ini MEMENUHI atau TIDAK MEMENUHI semua syarat di atas. 
                Sebutkan dengan jelas bagian mana yang tidak sesuai atau salah jika kandidat gagal memenuhi syarat.
                
                Berikan jawaban dalam format singkat berikut:
                - Status: [LULUS / TIDAK LULUS]
                - Alasan / Bagian yang Tidak Sesuai (jika ada yang melanggar jouken):
                - Catatan Penting:
                """
                
                # Panggil Gemini API (menggunakan model 3.5-flash)
                try:
                    response = client.models.generate_content(
                        model='gemini-3.5-flash',
                        contents=prompt
                    )
                    result_text = response.text
                except Exception as e:
                    result_text = f"Error saat memproses dengan Gemini: {str(e)}"
                
                # Tampilkan hasil per kandidat di web dalam bentuk kotak/card rapi
                with st.expander(f"Hasil Analisis: {current_sheet_name}", expanded=True):
                    st.markdown(result_text)
                
                # Update progress bar
                progress_bar.progress(i / total_candidates)
                
                # Jeda kecil agar aman dari limit
                time.sleep(2)
                
            st.success("Semua kandidat berhasil diperiksa!")
            
        except Exception as e:
            st.error(f"Terjadi kesalahan saat membaca file Excel: {str(e)}")
